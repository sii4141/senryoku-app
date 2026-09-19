import openpyxl
from openpyxl.utils.cell import column_index_from_string

path = r"C:\Users\SHUN\Downloads\新・戦力評価.xlsx"
workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
master = workbook["データ(変更禁止)"]
owned = workbook["所持設計図"]
target = "ワイルドファイア-格闘護送艦"

hits = []
master.reset_dimensions()
for row in range(2, 2000):
    name = str(master.cell(row, 19).value or "").strip()
    column = str(master.cell(row, 20).value or "").strip()
    if "ワイルドファイア" in name or name == target:
        hits.append((row, name, column))
    if row > 100 and not name and not column:
        break

print("sheets:", workbook.sheetnames)
print("matches:", hits)
print("owned size:", owned.calculate_dimension(force=True))
for _, name, column in hits:
    try:
        index = column_index_from_string(column)
        print(name, column, index, owned.cell(1, index).value, owned.cell(2, index).value)
    except Exception as error:
        print("invalid column:", column, error)
